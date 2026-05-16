from __future__ import annotations

import json
import re
import shutil
from collections import defaultdict
from pathlib import Path

import cv2
import numpy as np
import openpyxl
from rapidocr_onnxruntime import RapidOCR


ROOT = Path(__file__).resolve().parents[2]
FACES_ROOT = ROOT / "牌面"
PUBLIC_FACES_ROOT = ROOT / "h5-app" / "public" / "cards" / "faces"
PUBLIC_CARDS_PATH = ROOT / "h5-app" / "public" / "data" / "cards.json"
DIST_CARDS_PATH = ROOT / "h5-app" / "dist" / "data" / "cards.json"
RENAME_MAP_PATH = ROOT / "asset-pipeline" / "data" / "processed" / "card-face-rename-map.json"
OCR_OUTPUT_PATH = ROOT / "asset-pipeline" / "data" / "processed" / "card-face-ocr.json"
ASSET_VERSION_PATH = ROOT / "h5-app" / "src" / "utils" / "assetVersion.ts"

CODE_RE = re.compile(r"S\d{2}-[A-Z0-9]+", re.IGNORECASE)
DISPLAY_SUFFIX_RE = re.compile(r"-\d+$")
DIGITS_RE = re.compile(r"\d+")
S2_WORKBOOK_GLOB = "*彼界*奥林匹斯*.xlsx"
S15_WORKBOOK_GLOB = "*四阵营*补强*.xlsx"

FOLDER_ALIAS = {
    "天庭": "天廷",
}

DISPLAY_NAME_OVERRIDES = {
    ("奥林匹斯", "士气"): "奥林匹斯/士气卡",
    ("奥林匹斯", "神力"): "奥林匹斯/神力卡",
    ("奥林匹斯", "奥林匹斯 诸神巅"): "奥林匹斯主城 诸神巅",
    ("彼界", "士气"): "彼界士气卡",
    ("彼界", "符文"): "彼界符文卡",
    ("彼界", "彼界 阿瓦隆"): "彼界主城 阿瓦隆",
}

SECTION_SERIES = {
    "奥林匹斯": "S2彼界&奥林匹斯",
    "彼界": "S2彼界&奥林匹斯",
    "天灾": "独立天灾图鉴",
}


def normalize_card_id(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "-", normalize_text(value).lower()).strip("-")
    return normalized or "card"


def normalize_text(value: str) -> str:
    return str(value).replace(" ", "").replace("\n", "").replace("·", "").replace("•", "").strip()


def display_name_from_file(image_path: Path) -> str:
    return DISPLAY_SUFFIX_RE.sub("", image_path.stem).strip() or image_path.stem


def natural_key(path: Path) -> list[object]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", path.name)]


def parse_int(value: object) -> int | None:
    text = str(value or "").strip()
    if not text or text == "/":
        return None
    match = DIGITS_RE.search(text)
    return int(match.group(0)) if match else None


def parse_health(value: object) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    match = DIGITS_RE.search(text)
    return int(match.group(0)) if match else None


def infer_type_from_attr(attr: str, folder: str, name: str) -> str:
    text = attr.strip()
    if text:
        return text
    if folder == "天灾":
        return "联动天灾" if name == "群鸦盛宴之梦" else "天灾终局"
    return ""


def infer_type_from_name(folder: str, name: str, code: str, current_type: str) -> str:
    if current_type:
        return current_type
    if code.endswith("D1"):
        return "主城"
    if code.endswith("C1") or code.endswith("C1A"):
        return "士气卡"
    if "神力" in name:
        return "士气卡"
    if "符文" in name:
        return "士气卡"
    if "主城" in name or name.endswith("诸神巅") or name.endswith("阿瓦隆") or name.endswith("黄泉之门") or name.endswith("灵霄宝殿") or name.endswith("众神之乡") or name.endswith("英灵殿"):
        return "主城"
    if name in {"湖中仙女的馈赠", "寻找圣杯之旅", "芬尼亚传奇", "十字军东征"}:
        return "试炼卡"
    if name in {"王者之剑"}:
        return "衍生卡"
    return ""


def resolve_excel_name(folder: str, display_name: str) -> str:
    return DISPLAY_NAME_OVERRIDES.get((folder, display_name), display_name)


def read_json(path: Path) -> list[dict]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_legacy_cards() -> list[dict]:
    source = DIST_CARDS_PATH if DIST_CARDS_PATH.exists() else PUBLIC_CARDS_PATH
    return read_json(source)


def build_legacy_indexes(cards: list[dict]) -> tuple[dict[str, dict], dict[tuple[str, str], list[dict]]]:
    code_index: dict[str, dict] = {}
    name_index: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for card in cards:
        card_no = normalize_text(card.get("cardNo", "")).upper()
        if card_no:
            code_index[card_no] = card
        key = (str(card.get("faction", "")).strip(), normalize_text(card.get("name", "")))
        name_index[key].append(card)
    return code_index, name_index


def build_effect_lookup() -> dict[tuple[str, str], dict]:
    lookup: dict[tuple[str, str], dict] = {}

    s2_path = next(ROOT.glob(S2_WORKBOOK_GLOB))
    wb = openpyxl.load_workbook(s2_path, data_only=True)
    ws = wb["Sheet1"]
    section = "奥林匹斯"
    for row in ws.iter_rows(values_only=True):
        values = [str(value).strip() if value is not None else "" for value in row[:7]]
        name = values[0]
        if not name or name == "卡名":
            continue
        if name == "彼界士气卡":
            section = "彼界"
        lookup[(section, normalize_text(name))] = {
            "effectText": values[1],
            "cost": parse_int(values[2]),
            "attack": parse_int(values[3]),
            "health": parse_health(values[6]) if values[5] == "主宰" else None,
            "type": infer_type_from_attr(values[5], section, name),
        }
    wb.close()

    s15_path = next(ROOT.glob(S15_WORKBOOK_GLOB))
    wb = openpyxl.load_workbook(s15_path, data_only=True)
    ws = wb["Sheet1"]
    section = ""
    for row in ws.iter_rows(values_only=True):
        values = [str(value).strip() if value is not None else "" for value in row[:7]]
        first = values[0]
        if first in {"阿斯加德", "高天原", "太阳城", "天庭", "通用", "天灾"}:
            section = FOLDER_ALIAS.get(first, first)
            continue
        if not first or first in {"卡名", "天灾名称"}:
            continue
        if not section:
            continue
        lookup[(section, normalize_text(first))] = {
            "effectText": values[1],
            "cost": parse_int(values[2]),
            "attack": parse_int(values[3]),
            "health": parse_health(values[6]) if values[5] == "主宰" else None,
            "type": infer_type_from_attr(values[5], section, first),
        }
    wb.close()
    return lookup


def detect_code(image_path: Path, ocr: RapidOCR) -> str:
    image = cv2.imdecode(np.fromfile(str(image_path), dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise RuntimeError(f"无法读取图片: {image_path}")
    height, width = image.shape[:2]
    crop = image[int(height * 0.84) : height, : int(width * 0.36)]
    result, _ = ocr(crop)
    merged = " ".join(str(item[1]).strip() for item in result or [] if str(item[1]).strip())
    match = CODE_RE.search(merged)
    return match.group(0).upper() if match else ""


def infer_series(folder: str, legacy_card: dict | None) -> str:
    if legacy_card and legacy_card.get("series"):
        return legacy_card["series"]
    return SECTION_SERIES.get(folder, folder)


def build_card(
    folder: str,
    filename: str,
    display_name: str,
    code: str,
    legacy_card: dict | None,
    effect_entry: dict | None,
) -> dict:
    if legacy_card:
        card = dict(legacy_card)
    else:
        fallback_no = code or display_name
        card = {
            "id": f"face-{normalize_text(fallback_no).lower().replace('-', '_')}",
            "alias": [],
            "series": infer_series(folder, None),
            "cardNo": fallback_no,
            "rarity": "",
            "cost": None,
            "attack": None,
            "health": None,
            "tags": [],
            "faqIds": [],
            "flavorText": "",
        }

    card["name"] = display_name
    card["image"] = f"/cards/faces/{folder}/{filename}"
    card["faction"] = folder
    card["series"] = infer_series(folder, legacy_card)
    card["cardNo"] = card.get("cardNo") or code or display_name
    card["id"] = normalize_card_id(card["cardNo"])
    card["verified"] = True
    source = dict(card.get("source", {}))
    source["file"] = f"牌面/{folder}/{filename}"
    source["code"] = code or source.get("code", "")
    card["source"] = source

    if effect_entry:
        if effect_entry["type"]:
            card["type"] = effect_entry["type"]
        card["effectText"] = effect_entry["effectText"] or card.get("effectText", "")
        card["cost"] = effect_entry["cost"]
        card["attack"] = effect_entry["attack"]
        if effect_entry["health"] is not None:
            card["health"] = effect_entry["health"]

    card["subType"] = card.get("subType", "")
    card["searchText"] = f"{card['name']} {card['faction']} {card.get('type', '')} {card.get('effectText', '')}".strip()
    return card


def main() -> None:
    if not FACES_ROOT.exists():
        raise RuntimeError("未找到牌面目录。")

    ocr = RapidOCR()
    legacy_cards = load_legacy_cards()
    legacy_code_index, legacy_name_index = build_legacy_indexes(legacy_cards)
    effect_lookup = build_effect_lookup()

    results: list[dict] = []
    rename_map: list[dict] = []
    ocr_map: list[dict] = []

    if PUBLIC_FACES_ROOT.exists():
        shutil.rmtree(PUBLIC_FACES_ROOT)

    for folder_path in sorted([path for path in FACES_ROOT.iterdir() if path.is_dir()], key=lambda item: item.name):
        folder = folder_path.name
        target_dir = PUBLIC_FACES_ROOT / folder
        target_dir.mkdir(parents=True, exist_ok=True)
        for image_path in sorted([path for path in folder_path.iterdir() if path.suffix.lower() == ".png"], key=natural_key):
            shutil.copy2(image_path, target_dir / image_path.name)
            display_name = display_name_from_file(image_path)
            excel_name = resolve_excel_name(folder, display_name)
            code = detect_code(image_path, ocr)

            legacy_card = legacy_code_index.get(normalize_text(code).upper()) if code else None
            if legacy_card is None:
                candidates = legacy_name_index.get((folder, normalize_text(excel_name)), [])
                if len(candidates) == 1:
                    legacy_card = candidates[0]

            effect_entry = effect_lookup.get((folder, normalize_text(excel_name)))
            card = build_card(folder, image_path.name, display_name, code, legacy_card, effect_entry)
            card["type"] = infer_type_from_name(folder, display_name, code, card.get("type", ""))
            card["searchText"] = f"{card['name']} {card['faction']} {card.get('type', '')} {card.get('effectText', '')}".strip()
            results.append(card)

            rename_map.append(
                {
                    "folder": folder,
                    "original": image_path.name,
                    "renamed": image_path.name,
                    "name": display_name,
                    "code": code,
                    "type": card.get("type", ""),
                    "legacyMatch": legacy_card.get("name") if legacy_card else None,
                }
            )
            ocr_map.append(
                {
                    "folder": folder,
                    "file": image_path.name,
                    "name": display_name,
                    "code": code,
                    "excelName": excel_name,
                }
            )

    results.sort(key=lambda item: (item.get("faction", ""), item.get("name", ""), item.get("image", "")))
    PUBLIC_CARDS_PATH.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    RENAME_MAP_PATH.write_text(json.dumps(rename_map, ensure_ascii=False, indent=2), encoding="utf-8")
    OCR_OUTPUT_PATH.write_text(json.dumps(ocr_map, ensure_ascii=False, indent=2), encoding="utf-8")

    version_text = ASSET_VERSION_PATH.read_text(encoding="utf-8")
    version_text = re.sub(r'"\d{8}-\d{2}"', '"20260503-04"', version_text)
    ASSET_VERSION_PATH.write_text(version_text, encoding="utf-8")

    print(f"rebuilt {len(results)} cards from verified face filenames")


if __name__ == "__main__":
    main()
